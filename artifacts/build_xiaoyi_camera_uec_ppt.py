#!/usr/bin/env python3
"""汇报 PPT：小艺 Camera Tab（UEC）预览流 / 稳定帧 / 深度方案（领导评审版）。

视觉方案：白底 + 靛蓝主色，按进程/职责着色的泳道；架构图与流程图使用真实连接线
（带箭头、虚线 IPC 边界、分叉总线），拍摄流程用 UML 时序图表达。
"""

import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------- 调色 ----------
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
PANEL = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ACCENT = RGBColor(0x4F, 0x46, 0xE5)        # 靛蓝：标题/强调
ACCENT_SOFT = RGBColor(0xEE, 0xF2, 0xFF)

CAM = RGBColor(0x25, 0x63, 0xEB)           # 相机 App / 框架 / HAL
CAM_SOFT = RGBColor(0xDB, 0xEA, 0xFE)
CAM_PANEL = RGBColor(0xF0, 0xF6, 0xFF)

XY = RGBColor(0x05, 0x96, 0x69)            # 小艺 UEC
XY_SOFT = RGBColor(0xD1, 0xFA, 0xE5)
XY_PANEL = RGBColor(0xF0, 0xFD, 0xF6)

DEP = RGBColor(0xD9, 0x77, 0x06)           # 深度 / 拍摄旁路
DEP_SOFT = RGBColor(0xFE, 0xF3, 0xC7)

RISK = RGBColor(0xE1, 0x1D, 0x48)
RISK_SOFT = RGBColor(0xFF, 0xE4, 0xE6)
WARN = RGBColor(0xCA, 0x8A, 0x04)
WARN_SOFT = RGBColor(0xFE, 0xF9, 0xC3)

GREY = RGBColor(0x94, 0xA3, 0xB8)
GREY_SOFT = RGBColor(0xF1, 0xF5, 0xF9)

FONT = "Microsoft YaHei"
TOTAL = 11
SW, SH = 13.333, 7.5


# ---------- 基础工具 ----------
def _style_run(run, size, bold, color):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", FONT)


def text(slide, x, y, w, h, content, size=12, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=4):
    """多行文本框；content 可为 str 或 [(str, size, bold, color), ...] / [str, ...]。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(spacing)
        if isinstance(line, tuple):
            s, sz, b, c = (line + (size, bold, color))[:4]
            r = p.add_run()
            r.text = s
            _style_run(r, sz, b, c)
        else:
            r = p.add_run()
            r.text = line
            _style_run(r, size, bold, color)
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, dashed=False, radius=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dashed:
            ln = s.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def rrect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, dashed=False, radius=0.12):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, line_w, dashed, radius)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, dashed=False):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, line_w, dashed)


def fill_text(s, lines, size=12, bold=False, color=INK, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, spacing=2):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(5)
    tf.margin_top = tf.margin_bottom = Pt(3)
    lines = lines if isinstance(lines, list) else [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(spacing)
        if isinstance(line, tuple):
            s_, sz, b, c = (line + (size, bold, color))[:4]
            r = p.add_run()
            r.text = s_
            _style_run(r, sz, b, c)
        else:
            r = p.add_run()
            r.text = line
            _style_run(r, size, bold, color)
    return s


def node(slide, x, y, w, h, title, sub=None, color=CAM, soft=CAM_SOFT,
         title_size=12, sub_size=9.5, outline=True, fill=None):
    """带色条的节点：白底/浅底 + 左侧色条 + 标题/副标题。"""
    s = rrect(slide, x, y, w, h, fill=fill or WHITE, line=color if outline else None, line_w=1.0)
    rect(slide, x + 0.03, y + 0.12, 0.06, h - 0.24, fill=color)
    lines = [(title, title_size, True, INK)]
    if sub:
        for ln in sub.split("\n"):
            lines.append((ln, sub_size, False, MUTED))
    tb = text(slide, x + 0.14, y, w - 0.2, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=1)
    return s


def pill(slide, x, y, w, h, label, color, size=10, text_color=WHITE):
    s = rrect(slide, x, y, w, h, fill=color, radius=0.5)
    fill_text(s, label, size=size, bold=True, color=text_color)
    return s


def badge(slide, cx, cy, n, color=ACCENT, r=0.17, size=10):
    s = shape(slide, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill=color)
    fill_text(s, str(n), size=size, bold=True, color=WHITE)
    s.text_frame.margin_left = s.text_frame.margin_right = Pt(0)
    return s


def connector(slide, x1, y1, x2, y2, color=GREY, width=1.25, head=True, tail=False,
              dashed=False, elbow=False):
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    c = slide.shapes.add_connector(kind, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle")
        e.set("w", "med")
        e.set("len", "med")
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle")
        e.set("w", "med")
        e.set("len", "med")
    return c


def edge_label(slide, cx, cy, w, label, color=MUTED, size=9, bg=WHITE):
    s = rrect(slide, cx - w / 2, cy - 0.13, w, 0.26, fill=bg, line=None, radius=0.3)
    fill_text(s, label, size=size, bold=False, color=color)
    s.text_frame.margin_top = s.text_frame.margin_bottom = Pt(0)
    return s


def diamond(slide, cx, cy, w, h, label, color=ACCENT, soft=ACCENT_SOFT, size=10):
    s = shape(slide, MSO_SHAPE.DIAMOND, cx - w / 2, cy - h / 2, w, h, fill=soft, line=color, line_w=1.25)
    fill_text(s, label, size=size, bold=True, color=INK)
    return s


def panel(slide, x, y, w, h, title=None, color=ACCENT, fill=PANEL, dashed=False, title_size=11):
    p = rrect(slide, x, y, w, h, fill=fill, line=LINE, line_w=0.75, dashed=dashed, radius=0.06)
    if title:
        pill(slide, x + 0.18, y - 0.16, max(1.4, 0.16 * len(title) + 0.5), 0.32, title, color, size=title_size)
    return p


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    return s


def header(slide, title, caption=None, page=None):
    rect(slide, 0.6, 0.42, 0.09, 0.62, fill=ACCENT)
    text(slide, 0.85, 0.34, 11.5, 0.45, title, size=24, bold=True, color=INK)
    if caption:
        text(slide, 0.85, 0.78, 11.8, 0.3, caption, size=12, color=MUTED)
    rect(slide, 0.6, 7.02, 12.13, 0.012, fill=LINE)
    text(slide, 0.6, 7.08, 6, 0.3, "小艺 Camera Tab（UEC）智能感知方案 · 评审", size=9, color=GREY)
    if page:
        text(slide, 11.7, 7.08, 1.03, 0.3, f"{page} / {TOTAL}", size=9, color=GREY, align=PP_ALIGN.RIGHT)


def legend(slide, x, y, items, size=9):
    cx = x
    for color, label in items:
        rect(slide, cx, y + 0.07, 0.22, 0.12, fill=color)
        text(slide, cx + 0.28, y, 1.6, 0.26, label, size=size, color=MUTED)
        cx += 0.28 + 0.12 * len(label) + 0.35


def bullets(slide, x, y, w, h, items, size=11, color=INK, spacing=5, marker="•", marker_color=ACCENT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(2)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i:
            p.space_before = Pt(spacing)
        r = p.add_run()
        r.text = marker + " "
        _style_run(r, size, True, marker_color)
        if isinstance(item, tuple):
            r1 = p.add_run()
            r1.text = item[0]
            _style_run(r1, size, True, color)
            r2 = p.add_run()
            r2.text = item[1]
            _style_run(r2, size, False, color)
        else:
            r1 = p.add_run()
            r1.text = item
            _style_run(r1, size, False, color)
    return tb


# ======================================================================
def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    # ------------------------------------------------------------ 1 封面
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, fill=WHITE)
    rect(s, 0, 0, 5.4, SH, fill=ACCENT)
    rect(s, 5.4, 0, 0.06, SH, fill=RGBColor(0x37, 0x30, 0xA3))
    text(s, 0.6, 1.5, 4.5, 0.4, "方案评审", size=13, bold=True, color=RGBColor(0xC7, 0xD2, 0xFE))
    text(s, 0.6, 2.0, 4.6, 2.2, [
        ("小艺 Camera Tab", 34, True, WHITE),
        ("UEC 智能感知方案", 34, True, WHITE),
    ], spacing=2)
    text(s, 0.6, 3.9, 4.5, 1.2, [
        ("预览流获取 · 稳定帧门控", 13, False, RGBColor(0xE0, 0xE7, 0xFF)),
        ("端侧模型推荐 · 拍摄取 RGB + 深度", 13, False, RGBColor(0xE0, 0xE7, 0xFF)),
    ], spacing=4)
    text(s, 0.6, 6.7, 4.5, 0.4, "小艺 × 相机 App × 相机框架 / HAL", size=10, color=RGBColor(0xC7, 0xD2, 0xFE))

    # 右侧：迷你链路示意
    cx0 = 6.3
    steps = [("相机 HAL", CAM, CAM_SOFT), ("相机 App", CAM, CAM_SOFT), ("小艺 UEC", XY, XY_SOFT),
             ("稳定帧门控", XY, XY_SOFT), ("端侧模型", XY, XY_SOFT), ("chips 执行", DEP, DEP_SOFT)]
    for i, (t, c, sf) in enumerate(steps):
        y = 1.25 + i * 0.86
        n = rrect(s, cx0 + 0.9, y, 3.6, 0.58, fill=sf, line=c, line_w=1.0, radius=0.3)
        fill_text(n, t, size=13, bold=True, color=INK)
        badge(s, cx0 + 0.5, y + 0.29, i + 1, color=c, r=0.16)
        if i < len(steps) - 1:
            connector(s, cx0 + 2.7, y + 0.58, cx0 + 2.7, y + 0.86, color=GREY, width=1.25)
    text(s, cx0 + 4.8, 1.3, 2.4, 3.5, [
        ("会话主人：相机 App", 11, True, INK),
        ("小艺不直连 HAL / HDI", 10, False, MUTED),
        ("", 6, False, MUTED),
        ("门控独立模块", 11, True, INK),
        ("稳定才推理，省电更准", 10, False, MUTED),
        ("", 6, False, MUTED),
        ("拍摄 = 一次带深度的 Capture", 11, True, INK),
        ("RGB / 深度 / meta 同 timestamp", 10, False, MUTED),
    ], spacing=3)

    # ------------------------------------------------------------ 2 背景与目标
    s = blank(prs)
    header(s, "背景与目标", "在相机内新增「小艺」Tab：看懂画面 → 主动推荐 → 一步执行", 2)

    cols = [
        ("用户价值", XY, XY_SOFT, ["取景时实时理解内容", "弹出可点击 chips，少操作直达结果", "识物 / 推荐 / AR 空间理解"]),
        ("业务价值", CAM, CAM_SOFT, ["小艺能力以 UEC 形式可插拔进相机", "独立迭代，不改相机主流程", "同一套感知能力可复用到其他入口"]),
        ("技术目标", DEP, DEP_SOFT, ["稳定获取预览流与运动状态", "按需获取对齐的 RGB + 深度", "端侧闭环：门控 → 模型 → 分发"]),
    ]
    for i, (t, c, sf, items) in enumerate(cols):
        x = 0.6 + i * 4.1
        rrect(s, x, 1.35, 3.9, 2.25, fill=WHITE, line=LINE, line_w=0.75)
        rect(s, x, 1.35, 3.9, 0.08, fill=c)
        text(s, x + 0.25, 1.55, 3.4, 0.4, t, size=15, bold=True, color=c)
        bullets(s, x + 0.25, 2.0, 3.45, 1.5, items, size=11, marker_color=c)

    panel(s, 0.6, 4.05, 12.13, 2.75, "核心能力链路", ACCENT, fill=PANEL)
    chain = [("预览流", "双 Surface\n显示 + 分析", CAM, CAM_SOFT), ("运动 meta", "随帧实时\n上传应用层", CAM, CAM_SOFT),
             ("稳定帧门控", "独立模块\n判定触发", XY, XY_SOFT), ("感知 + 端侧模型", "内容识别\n推荐分析", XY, XY_SOFT),
             ("分发 / chips", "规则弹出\n点击执行", XY, XY_SOFT), ("拍摄 RGB+深度", "按需单次\n时间戳对齐", DEP, DEP_SOFT)]
    for i, (t, d, c, sf) in enumerate(chain):
        x = 0.9 + i * 1.95
        n = rrect(s, x, 4.6, 1.6, 1.3, fill=sf, line=c, line_w=1.0)
        fill_text(n, [(t, 12, True, INK)] + [(ln, 9.5, False, MUTED) for ln in d.split("\n")], spacing=1)
        badge(s, x + 0.8, 4.6, i + 1, color=c, r=0.15, size=9)
        if i < len(chain) - 1:
            connector(s, x + 1.6, 5.25, x + 1.95, 5.25, color=GREY)
    text(s, 0.9, 6.15, 11.6, 0.4,
         "原则：摄像头会话只有一个主人（相机 App）；小艺是流的消费者与业务决策者，不直连 HAL / HDI。",
         size=11, bold=True, color=ACCENT)

    # ------------------------------------------------------------ 3 总体架构
    s = blank(prs)
    header(s, "总体架构", "相机进程握相机；小艺 UEC 进程消费流并做感知、推荐与 AR", 3)

    # 相机进程容器
    panel(s, 0.6, 1.5, 5.3, 5.2, "相机 App 进程（宿主）", CAM, fill=CAM_PANEL)
    lx, lw = 0.95, 4.6
    n1 = node(s, lx, 1.85, lw, 0.85, "相机业务 + 对接小艺接口层", "Start / Stop 预览 · CaptureAt · 抢占通知", CAM, CAM_SOFT)
    n2 = node(s, lx, 3.0, lw, 0.75, "Camera Framework", "Session / Output 管理，下发 Capture", CAM, CAM_SOFT)
    n3 = rrect(s, lx, 4.05, lw, 0.5, fill=GREY_SOFT, line=GREY, line_w=0.75, dashed=True)
    fill_text(n3, "HDI ─ 框架 ↔ HAL 硬件接口层（应用不直连）", size=10, bold=True, color=MUTED)
    n4 = node(s, lx, 4.85, lw, 1.55, "Camera HAL · Sensor / ISP",
              "同一帧投两路 Output：Surface A（显示→RS）/ Surface B（分析→小艺）\n连续：预览 Buffer + 运动 meta（同 ts）；按需：单次深度 + 对齐 RGB", CAM, CAM_SOFT)
    mid = lx + lw / 2
    connector(s, mid, 2.7, mid, 3.0, color=CAM, head=True, tail=True)
    connector(s, mid, 3.75, mid, 4.05, color=CAM, head=True, tail=True)
    connector(s, mid, 4.55, mid, 4.85, color=CAM, head=True, tail=True)

    # IPC 边界
    connector(s, 6.65, 1.45, 6.65, 6.75, color=GREY, width=1.25, head=False, dashed=True)

    # 小艺进程容器
    panel(s, 7.4, 1.5, 5.33, 5.2, "小艺 UEC 进程（Tab）", XY, fill=XY_PANEL)
    rx = 7.7
    # 两路 Surface 入口
    sa = rrect(s, rx, 1.85, 1.65, 0.7, fill=GREY_SOFT, line=GREY, line_w=1.0)
    fill_text(sa, [("Surface A · 显示路", 10, True, INK), ("消费者 = RS 直达屏幕", 8, False, MUTED)], spacing=0)
    sb = rrect(s, rx + 1.75, 1.85, 3.0, 0.7, fill=XY_SOFT, line=XY, line_w=1.0)
    fill_text(sb, [("Surface B · 分析路（ImageReceiver）", 10, True, INK),
                   ("消费者 = 小艺；小图低帧率；acquire→处理→release", 8, False, MUTED)], spacing=0)
    # 显示直下
    connector(s, rx + 1.45 / 2, 2.55, rx + 1.45 / 2, 3.15, color=GREY)
    # 分析路分叉总线
    bus_y = 2.85
    connector(s, rx + 1.75 + 3.0 / 2, 2.55, rx + 1.75 + 3.0 / 2, bus_y, color=XY, head=False)
    connector(s, rx + 1.6 + 1.55 / 2, bus_y, rx + 3.3 + 1.45 / 2, bus_y, color=XY, head=False)
    kids = [("显示 XComponent", "宿主还是 UEC：争议 G", rx, 1.45, GREY),
            ("稳定帧门控", "独立模块 · motion meta", rx + 1.6, 1.55, XY),
            ("AR 空间建模", "共用 B 或第三路 Output", rx + 3.3, 1.45, XY)]
    for t, d, kx, kw, c in kids:
        if c is XY:
            connector(s, kx + kw / 2, bus_y, kx + kw / 2, 3.15, color=XY)
        n = rrect(s, kx, 3.15, kw, 0.75, fill=WHITE, line=c, line_w=1.0)
        fill_text(n, [(t, 10.5, True, INK), (d, 8.5, False, MUTED)], spacing=0)
    # 门控 → 感知/模型 → 分发
    connector(s, rx + 1.6 + 1.55 / 2, 3.9, rx + 1.6 + 1.55 / 2, 4.2, color=XY)
    n = rrect(s, rx + 1.1, 4.2, 2.55, 0.65, fill=WHITE, line=XY, line_w=1.0)
    fill_text(n, [("感知 / 端侧大模型", 11, True, INK), ("仅在稳定帧触发；异步、过期丢弃", 8.5, False, MUTED)], spacing=0)
    connector(s, rx + 1.1 + 2.55 / 2, 4.85, rx + 1.1 + 2.55 / 2, 5.15, color=XY)
    n = rrect(s, rx + 1.1, 5.15, 2.55, 0.6, fill=WHITE, line=XY, line_w=1.0)
    fill_text(n, [("分发服务 + 规则 → chips", 11, True, INK), ("冷却 / 去重 / 点击执行", 8.5, False, MUTED)], spacing=0)
    # 拍摄旁路
    cap = rrect(s, rx, 5.95, 4.75, 0.6, fill=DEP_SOFT, line=DEP, line_w=1.0)
    fill_text(cap, [("拍摄按钮 → CaptureAt(ts) → RGB + 深度(fd) + meta", 10.5, True, INK),
                    ("对齐校验后送模型 / AR；显示层 XComponent 可用性待实测", 8.5, False, MUTED)], spacing=0)

    # 跨进程连线：预览流两路（接口层 → 小艺；显示路为灰、分析路为蓝）
    connector(s, lx + lw, 1.95, rx, 1.95, color=GREY, width=1.5)
    connector(s, lx + lw, 2.15, rx + 1.75, 2.15, color=CAM, width=1.75)
    edge_label(s, 6.65, 1.68, 2.2, "A 显示 / B 分析 + meta", color=CAM, bg=WHITE)
    # 拍摄请求：小艺拍摄框 → 接口层（所有跨进程调用都进接口层，不直达 HAL）
    connector(s, rx, 6.2, 6.3, 6.2, color=DEP, width=1.5, head=False)
    connector(s, 6.3, 6.2, 6.3, 2.3, color=DEP, width=1.5, head=False)
    connector(s, 6.3, 2.3, lx + lw, 2.3, color=DEP, width=1.5)
    # 结果回传：接口层 → 小艺（虚线）
    connector(s, lx + lw, 2.55, 7.0, 2.55, color=DEP, width=1.5, head=False, dashed=True)
    connector(s, 7.0, 2.55, 7.0, 6.42, color=DEP, width=1.5, head=False, dashed=True)
    connector(s, 7.0, 6.42, rx, 6.42, color=DEP, width=1.5, dashed=True)
    lab = rrect(s, 6.15, 3.3, 1.0, 0.5, fill=WHITE, line=GREY, line_w=0.75, radius=0.3)
    fill_text(lab, [("进程边界", 9, True, MUTED), ("IPC", 9, False, MUTED)], spacing=0)
    lab2 = rrect(s, 6.0, 4.6, 1.3, 0.85, fill=WHITE, line=DEP, line_w=0.75, radius=0.15)
    fill_text(lab2, [("── 实线", 8, True, DEP), ("CaptureAt(ts)", 8, False, INK),
                     ("╌╌ 虚线", 8, True, DEP), ("RGB+Depth fd+meta", 8, False, INK)], spacing=0)

    # ------------------------------------------------------------ 4 关键设计①
    s = blank(prs)
    header(s, "关键设计 ①｜预览流获取：显示 + 分析双 Surface", "一个 Surface 只有一个消费者；显示路被 RS 直接消费，应用层截不到帧 → 必须两路 Output", 4)

    flow = [("进 Tab", "UEC 拉起", GREY, GREY_SOFT), ("交换两路 Surface", "A：显示（XComponent）\nB：分析（ImageReceiver）", XY, XY_SOFT),
            ("相机配流", "Session 加两路 Output\n各自 size / fps", CAM, CAM_SOFT), ("HAL 出图", "同一帧投 A、B\n运动 meta 同 ts", CAM, CAM_SOFT),
            ("小艺消费 B", "门控 / 模型 / AR", XY, XY_SOFT)]
    for i, (t, d, c, sf) in enumerate(flow):
        x = 0.7 + i * 2.45
        n = rrect(s, x, 1.4, 2.0, 1.1, fill=sf, line=c, line_w=1.0)
        fill_text(n, [(t, 12.5, True, INK)] + [(ln, 9.5, False, MUTED) for ln in d.split("\n")], spacing=1)
        badge(s, x + 1.0, 1.4, i + 1, color=c, r=0.15, size=9)
        if i < len(flow) - 1:
            connector(s, x + 2.0, 1.95, x + 2.45, 1.95, color=GREY)

    # 左：双 Surface 图（生产者 → 队列 → 唯一消费者）
    panel(s, 0.6, 3.0, 6.6, 3.8, "双 Surface：为什么不能单路 fork（B 已收敛）", ACCENT, fill=PANEL)
    src = rrect(s, 0.85, 4.35, 1.5, 1.0, fill=CAM_SOFT, line=CAM, line_w=1.0)
    fill_text(src, [("Camera HAL", 11.5, True, INK), ("生产者", 9, False, MUTED), ("同一帧投两路", 8.5, False, MUTED)], spacing=0)
    connector(s, 2.35, 4.85, 2.7, 4.85, color=CAM, head=False, width=1.5)
    connector(s, 2.7, 3.85, 2.7, 5.85, color=CAM, head=False, width=1.5)
    lanes4 = [
        (3.85, "Surface A", "显示路 · 全分辨率 30fps", GREY, GREY_SOFT, "RS 合成上屏", "应用层不在队列上"),
        (5.85, "Surface B", "分析路 · 小图 1～5fps", XY, XY_SOFT, "小艺 ImageReceiver", "acquire → 处理 → release"),
    ]
    for yy, t, d, c, sf, cons, cd in lanes4:
        connector(s, 2.7, yy, 3.05, yy, color=CAM, width=1.5)
        q = rrect(s, 3.05, yy - 0.38, 1.75, 0.76, fill=sf, line=c, line_w=1.0)
        fill_text(q, [(t, 11, True, INK), (d, 8.5, False, MUTED)], spacing=0)
        connector(s, 4.8, yy, 5.15, yy, color=c, width=1.5)
        k = rrect(s, 5.15, yy - 0.38, 1.85, 0.76, fill=WHITE, line=c, line_w=1.0)
        fill_text(k, [(cons, 10.5, True, INK), (cd, 8.5, False, MUTED)], spacing=0)
    text(s, 3.05, 4.62, 3.95, 0.5, [("规则：一个 Surface 只有一个消费者", 9.5, True, ACCENT),
                                    ("HAL flush → RS 立刻 acquire → 上屏，小艺没有插手机会", 8.5, False, MUTED)], spacing=0, align=PP_ALIGN.CENTER)
    text(s, 0.85, 6.3, 6.1, 0.45, "不做：小艺先消费 B 再转发到显示（多一次拷贝 + ≥1 帧延迟）；AR 尺寸不同再加第三路 Output",
         size=9.5, bold=True, color=RISK)

    # 右：契约
    panel(s, 7.45, 3.0, 5.28, 3.8, "格式与元数据契约", DEP, fill=PANEL)
    bullets(s, 7.7, 3.35, 4.9, 3.3, [
        ("BT709_FULL", "：BT.709 色域语义 + Full range，规定「数值怎么解释」"),
        ("需写死", "：实际像素格式（RGB / NV12…）、位深、stride、旋转责任方"),
        ("运动 meta", "：随预览每帧上传 gyro / accel / 对焦 / 曝光"),
        ("硬约束", "：meta.timestamp 与 frame.timestamp 可对齐；meta 缺失不得判稳"),
        ("Surface 归属（A）", "：A 由 XComponent 所在进程建、B 由小艺建；切 Tab 时谁先停"),
        ("能力协商", "：HAL 最大并发 Output 数、各路允许的分辨率组合"),
    ], size=10.5, spacing=6, marker_color=DEP)

    # ------------------------------------------------------------ 5 关键设计②
    s = blank(prs)
    header(s, "关键设计 ②｜稳定帧门控 → 感知 → 模型", "门控独立成模块，挂在 Tab 业务层；感知只消费稳定事件", 5)

    y0 = 2.05
    a = rrect(s, 0.7, y0 - 0.45, 1.9, 0.9, fill=CAM_SOFT, line=CAM, line_w=1.0)
    fill_text(a, [("预览帧 + 运动 meta", 11, True, INK), ("每帧，很轻", 9, False, MUTED)], spacing=0)
    connector(s, 2.6, y0, 3.05, y0, color=GREY)
    g = rrect(s, 3.05, y0 - 0.55, 2.3, 1.1, fill=XY_SOFT, line=XY, line_w=1.5)
    fill_text(g, [("StableFrameDetector", 11, True, INK), ("|gyro| / accel 阈值", 9, False, MUTED), ("连续 K 帧 / 持续 M ms", 9, False, MUTED)], spacing=0)
    connector(s, 5.35, y0, 5.8, y0, color=GREY)
    diamond(s, 6.45, y0, 1.3, 1.0, "稳定？", XY, XY_SOFT, size=11)
    connector(s, 7.1, y0, 7.55, y0, color=XY)
    edge_label(s, 7.32, y0 - 0.3, 0.5, "是", color=XY)
    p1 = rrect(s, 7.55, y0 - 0.45, 1.7, 0.9, fill=WHITE, line=XY, line_w=1.0)
    fill_text(p1, [("感知 Perception", 11, True, INK), ("检测 / 识别", 9, False, MUTED)], spacing=0)
    connector(s, 9.25, y0, 9.65, y0, color=GREY)
    p2 = rrect(s, 9.65, y0 - 0.45, 1.6, 0.9, fill=WHITE, line=XY, line_w=1.0)
    fill_text(p2, [("端侧大模型", 11, True, INK), ("异步 · 过期丢弃", 9, False, MUTED)], spacing=0)
    connector(s, 11.25, y0, 11.6, y0, color=GREY)
    p3 = rrect(s, 11.6, y0 - 0.45, 1.13, 0.9, fill=DEP_SOFT, line=DEP, line_w=1.0)
    fill_text(p3, [("chips", 11, True, INK), ("规则 / 冷却", 9, False, MUTED)], spacing=0)
    # 否 分支
    connector(s, 6.45, y0 + 0.5, 6.45, 3.15, color=RISK)
    edge_label(s, 6.75, y0 + 0.75, 0.5, "否", color=RISK)
    dn = rrect(s, 5.55, 3.15, 1.8, 0.5, fill=RISK_SOFT, line=RISK, line_w=1.0)
    fill_text(dn, "丢弃 / 取消进行中推理", size=9.5, bold=True, color=RISK)
    # 订阅同一事件
    connector(s, 4.2, y0 + 0.55, 4.2, 3.4, color=XY, head=False, dashed=True)
    connector(s, 4.2, 3.4, 2.0, 3.4, color=XY, head=False, dashed=True)
    sub = rrect(s, 0.7, 3.15, 1.3, 0.5, fill=WHITE, line=XY, line_w=0.75, dashed=True)
    fill_text(sub, "AR 关键帧 / 深度\n同一事件订阅", size=8.5, color=MUTED, spacing=0)

    panel(s, 0.6, 4.15, 6.0, 2.65, "为什么独立成模块（而非塞进感知前）", XY, fill=PANEL)
    rows = [("职责", "稳不稳（传感器）vs 有什么（内容），关注点分离"),
            ("复用", "模型、AR 关键帧、深度、chips 都订同一事件"),
            ("节奏", "每帧轻量先挡掉大部分重活，可单独调参 / 关闭"),
            ("测试", "喂假 meta 即可单测，不依赖整条感知链")]
    for i, (k, v) in enumerate(rows):
        yy = 4.5 + i * 0.55
        pill(s, 0.85, yy, 0.8, 0.36, k, XY, size=10)
        text(s, 1.8, yy, 4.6, 0.4, v, size=10.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    panel(s, 6.85, 4.15, 5.88, 2.65, "门控输入 / 输出 / 参数", ACCENT, fill=PANEL)
    bullets(s, 7.1, 4.5, 5.45, 2.2, [
        ("输入", "：MotionMetadata（gyro / accel / afState / exposure，同 ts）"),
        ("辅助", "：轻量帧差 / 光流，用于无 meta 兜底"),
        ("输出", "：OnStableFrame(ts)、OnUnstable()"),
        ("参数（L）", "：阈值、K 帧、M ms、冷却时间"),
        ("说明", "：点击拍摄仍可强制 Capture，不必非稳不可（产品定）"),
    ], size=10.5, spacing=5)

    # ------------------------------------------------------------ 6 关键设计③ 时序图
    s = blank(prs)
    header(s, "关键设计 ③｜点击拍摄：RGB + 深度 + 运动 meta", "不是从旧预览里「挖」深度，而是按 timestamp 触发一次带深度的 Capture", 6)

    lanes = [("用户", GREY, GREY_SOFT, 1.3), ("小艺 UEC", XY, XY_SOFT, 3.9), ("相机 App", CAM, CAM_SOFT, 6.7), ("框架 → HDI → HAL", CAM, CAM_SOFT, 9.5)]
    top_y, bot_y = 1.45, 6.75
    for name, c, sf, cx in lanes:
        hd = rrect(s, cx - 0.95, top_y, 1.9, 0.5, fill=sf, line=c, line_w=1.0)
        fill_text(hd, name, size=11, bold=True, color=INK)
        connector(s, cx, top_y + 0.5, cx, bot_y, color=LINE, width=1.0, head=False, dashed=True)

    def msg(y, a, b, label, color, dashed=False, width=1.25):
        xa, xb = lanes[a][3], lanes[b][3]
        connector(s, xa, y, xb, y, color=color, width=width, dashed=dashed)
        text(s, min(xa, xb) + 0.1, y - 0.3, abs(xb - xa) - 0.2, 0.28, label, size=9,
             color=color if dashed else INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)

    def note(y, cx, label, color=DEP, soft=DEP_SOFT, w=2.2):
        n = rrect(s, cx - w / 2, y - 0.2, w, 0.4, fill=soft, line=color, line_w=0.75)
        fill_text(n, label, size=9, color=INK)

    msg(2.45, 0, 1, "① 点击拍摄（记 ts = T）", INK)
    msg(3.0, 1, 2, "② CaptureAt(T, rgb, depth, pause)", XY)
    msg(3.55, 2, 3, "③ 单次 Capture + 深度（可短暂停 · C）", CAM)
    note(4.05, 9.3, "HAL：抓近 T 的 RGB · 单次算深度 · 附运动 meta", CAM, CAM_SOFT, 3.1)
    msg(4.55, 3, 2, "④ RGB + Depth(fd) + meta @T · 经 HDI", CAM, dashed=True)
    msg(5.05, 2, 1, "⑤ OnCaptureResult(rgb, depthFd, meta)", XY, dashed=True)
    note(5.5, 3.9, "⑥ map fd 读深度 → 对齐校验 → 模型 / AR → 释放 fd", XY, XY_SOFT, 3.6)
    msg(5.95, 1, 2, "⑦ ResumePreview（若曾 Pause）", XY)
    msg(6.4, 1, 0, "⑧ chips / 推荐结果", DEP, dashed=True)

    # 右侧备注
    panel(s, 11.0, 1.45, 1.73, 5.3, None, fill=PANEL)
    text(s, 11.1, 1.55, 1.55, 5.1, [
        ("名词", 10, True, ACCENT),
        ("fd：跨进程共享内存的「取货号」", 8.5, False, INK),
        ("handle：指向 Buffer 的凭证", 8.5, False, INK),
        ("HDI：框架 ↔ HAL 接口层，小艺不直连", 8.5, False, INK),
        ("", 4, False, INK),
        ("对齐", 10, True, DEP),
        ("RGB / 深度 / meta 同 ts；对不齐 → TIMESTAMP_MISMATCH 不喂模型（H）", 8.5, False, INK),
        ("", 4, False, INK),
        ("待拍板", 10, True, RISK),
        ("C 停流否", 8.5, False, INK),
        ("D 预览帧 vs Still", 8.5, False, INK),
        ("K fd 所有权 / 释放", 8.5, False, INK),
    ], spacing=3)

    # ------------------------------------------------------------ 7 接口清单
    s = blank(prs)
    header(s, "接口清单（小艺 ↔ 相机 App）", "五组接口；命名可改，语义保留；应用不直接调 HDI", 7)
    groups = [
        ("能力协商", CAM, [("QueryCameraCapability", "→ sizes / formats / fps / 最大并发 Output 数 / motion meta / 按需深度 / 可否 Pause")]),
        ("预览会话", XY, [("ProvideSurfaces", "[{surfaceId, role: DISPLAY | ANALYSIS}] 两路（争议 A：各由谁建）"),
                       ("StartPreviewForXiaoyi", "outputs: [{surfaceId, size, format, fps, role}] · needMotionMeta"),
                       ("OnPreviewStarted / Stopped", "以实配为准；Stopped 后可拆 Surface"), ("Pause / ResumePreview", "拍摄取深度时短暂停")]),
        ("帧 + 运动 meta", XY, [("OnPreviewFrame", "或 Surface 隐式出帧：timestamp · buffer · transform"),
                            ("OnPreviewMetadata", "同 ts：gyro / accel / 对焦 / 曝光"), ("OnStableFrame(ts)", "小艺内部事件，非跨进程")]),
        ("拍摄 / 深度", DEP, [("CaptureAt", "targetTimestamp · needRgb · needDepth · allowPause · timeout"),
                          ("OnCaptureResult", "rgb · depth fd/handle · 宽高 / 单位 / 置信度 · meta"), ("OnCaptureFailed / CancelCapture", "code：TIMEOUT / DEPTH_UNAVAILABLE …")]),
        ("生命周期 / 抢占", GREY, [("OnXiaoyiTabShow / Hide", "Hide 停推理 / AR"), ("OnHostCameraBusy", "宿主拍照 / 录像要回主预览"), ("ReleaseAll", "双向；断链 / 进程异常")]),
    ]
    y = 1.3
    row_h, pad, gap = 0.27, 0.07, 0.08
    for name, c, items in groups:
        h = row_h * len(items) + pad * 2
        rect(s, 0.6, y, 0.08, h, fill=c)
        rrect(s, 0.68, y, 12.05, h, fill=WHITE, line=LINE, line_w=0.75, radius=0.04)
        text(s, 0.85, y, 1.7, h, name, size=11, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
        for i, (api, desc) in enumerate(items):
            yy = y + pad + i * row_h
            text(s, 2.6, yy, 3.4, row_h, api, size=10, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
            text(s, 6.0, yy, 6.6, row_h, desc, size=9.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        y += h + gap
    rrect(s, 0.6, y + 0.02, 12.13, 6.9 - y, fill=ACCENT_SOFT, line=None)
    text(s, 0.85, y + 0.08, 11.7, 0.3,
         "关键说明：即便「画面画在小艺 Tab 内」，只要 Session 仍在相机 App，预览会话这组接口仍必须保留（争议 F）。",
         size=10, bold=True, color=ACCENT)
    text(s, 0.85, y + 0.38, 11.7, 0.3,
         "错误码：OK · NO_PERMISSION · CAMERA_IN_USE · SURFACE_INVALID · FORMAT_UNSUPPORTED · TIMEOUT · DEPTH_UNAVAILABLE · TIMESTAMP_MISMATCH · DEVICE_THERMAL",
         size=9.5, color=MUTED)

    # ------------------------------------------------------------ 8 风险与应对
    s = blank(prs)
    header(s, "风险与应对", "最大风险不是 UEC 切 Tab 慢那一下，而是「流归属 + UEC 能力边界 + 实时负载」", 8)
    risks = [
        ("UEC 内起流 / 显示受限", "高", "官方能力表曾列 UEC 不支持 XComponent，显示链路可能不可用", "样机实测；备选：宿主画预览，小艺只吃分析流 + 叠 UI"),
        ("与主相机抢会话", "高", "小艺再开一套 Session 会打断主预览或失败", "单会话多输出 / Tab 级移交；进出 Tab 明确停启"),
        ("端侧 LLM 跟帧负载", "高", "满帧送模型 → 发热、掉帧、拖垮预览", "稳定帧门控 + 降频 + 异步队列 + 过期丢弃"),
        ("推荐迟到 / chips 乱弹", "中高", "检测 + LLM 冷路径 1～3s，场景已切换", "场景稳定再推；结果带 ts；冷却、去重、置信度门槛"),
        ("UEC 切换体感", "中", "首次进 Tab 冷启几百 ms～1s，闪白", "同色背景 + placeholder；评估进程保活策略"),
        ("生命周期 / 黑屏 / 泄漏", "中高", "Surface 销毁顺序、fd 未释放、后台未停流", "显隐 / 前后台统一停启；fd 所有权写进契约"),
    ]
    for i, (t, lvl, why, act) in enumerate(risks):
        col, row = divmod(i, 3)
        x = 0.6 + col * 6.15
        y = 1.4 + row * 1.82
        lc, ls = (RISK, RISK_SOFT) if lvl == "高" else (WARN, WARN_SOFT)
        rrect(s, x, y, 5.98, 1.65, fill=WHITE, line=LINE, line_w=0.75)
        rect(s, x, y + 0.15, 0.07, 1.35, fill=lc)
        pill(s, x + 0.25, y + 0.18, 0.6, 0.3, lvl, lc, size=9.5)
        text(s, x + 0.95, y + 0.14, 4.9, 0.38, t, size=13, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, y + 0.6, 5.6, 0.45, [("风险  ", 9.5, True, MUTED), (why, 10, False, INK)], spacing=0)
        tb = text(s, x + 0.25, y + 1.08, 5.6, 0.45, "", size=10)
        p = tb.text_frame.paragraphs[0]
        r = p.add_run(); r.text = "应对  "; _style_run(r, 9.5, True, ACCENT)
        r = p.add_run(); r.text = act; _style_run(r, 10, True, INK)

    # ------------------------------------------------------------ 9 待评审决策点
    s = blank(prs)
    header(s, "待评审决策点", "建议按顺序拍板：会话主人 → Surface / 显示 → 格式 / meta → 分发 / 门控 → 拍摄路径 → 二级跳转", 9)
    headers = ["ID", "议题", "选项", "建议 / 影响"]
    widths = [1.0, 2.5, 4.5, 4.13]
    x0 = 0.6
    rect(s, x0, 1.4, sum(widths), 0.42, fill=INK)
    xx = x0
    for i, h in enumerate(headers):
        text(s, xx + 0.1, 1.4, widths[i] - 0.2, 0.42, h, size=11, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        xx += widths[i]
    rows = [
        ("F/M", "会话主人 & 主预览策略", "相机 App 握机（小艺消费）vs 小艺自 openCamera", "建议相机握机；进 Tab 停主预览或双 Output", ACCENT),
        ("A/G", "Surface 归属 & UEC 显示", "小艺建 vs 相机建；UEC 能否 XComponent", "先实测 G；不可用则宿主显示", RISK),
        ("J/I", "格式 & 运动 meta", "RGB vs YUV；meta 字段集与每帧保障", "写死 pixelFormat / stride / 旋转；meta 同 ts", DEP),
        ("B/L", "多路 Output & 门控参数", "B 已收敛双 Surface；AR 是否第三路待定；门控阈值", "确认 HAL 并发 Output 上限；门控参数样机调", XY),
        ("C/D/H/K", "拍摄路径", "停流否；预览帧 vs Still；δt 补偿；fd 所有权", "允许短暂停；对齐失败不喂模型", DEP),
        ("E", "跳第二个相机 UEC", "宿主切换 vs UEC 内嵌套", "宿主切换 + 先停当前流", CAM),
    ]
    for ri, (rid, topic, opt, adv, c) in enumerate(rows):
        y = 1.82 + ri * 0.74
        rect(s, x0, y, sum(widths), 0.74, fill=PANEL if ri % 2 == 0 else WHITE)
        rect(s, x0, y + 0.74 - 0.01, sum(widths), 0.01, fill=LINE)
        pill(s, x0 + 0.12, y + 0.21, 0.76, 0.32, rid, c, size=9)
        xx = x0 + widths[0]
        for ci, cell in enumerate((topic, opt, adv)):
            text(s, xx + 0.1, y, widths[ci + 1] - 0.2, 0.74, cell, size=10.5, bold=(ci == 0),
                 color=INK if ci == 0 else MUTED, anchor=MSO_ANCHOR.MIDDLE)
            xx += widths[ci + 1]
    text(s, 0.6, 6.4, 12.1, 0.4, "另需产品 / 合规确认 N：预览与深度是否落盘、端侧模型是否出域、权限告知。",
         size=10.5, bold=True, color=WARN)

    # ------------------------------------------------------------ 10 推进与验证计划
    s = blank(prs)
    header(s, "推进与验证计划", "先用样机验证两道硬门，再细化规则与交互", 10)
    phases = [
        ("阶段 1", "技术预研（硬门）", ["双 Surface：显示路 + 小艺分析路并发出流", "UEC 内 XComponent 显示是否可用", "运动 meta 每帧上传与 ts 对齐"], RISK, RISK_SOFT),
        ("阶段 2", "闭环打通", ["稳定帧门控 → 感知 → 端侧模型", "chips 弹出与点击执行", "端到端时延 / 发热基线"], XY, XY_SOFT),
        ("阶段 3", "深度与拍摄", ["CaptureAt → RGB + 深度 fd 回传", "停流策略与对齐校验", "AR 建模接入"], DEP, DEP_SOFT),
        ("阶段 4", "体验打磨", ["UEC 冷启闪白消减", "chips 冷却 / 去重 / 阈值调优", "抢占、后台、异常恢复"], CAM, CAM_SOFT),
    ]
    # 时间轴
    connector(s, 0.9, 2.0, 12.6, 2.0, color=LINE, width=2.0, head=False)
    for i, (n, t, items, c, sf) in enumerate(phases):
        x = 0.6 + i * 3.05
        cx = x + 1.45
        shape(s, MSO_SHAPE.OVAL, cx - 0.2, 1.8, 0.4, 0.4, fill=c)
        fill_text(s.shapes[-1], str(i + 1), size=11, bold=True, color=WHITE)
        connector(s, cx, 2.2, cx, 2.55, color=c, head=False)
        rrect(s, x, 2.55, 2.9, 2.9, fill=WHITE, line=LINE, line_w=0.75)
        rect(s, x, 2.55, 2.9, 0.07, fill=c)
        pill(s, x + 0.2, 2.78, 0.9, 0.3, n, c, size=9.5)
        text(s, x + 0.2, 3.15, 2.5, 0.4, t, size=14, bold=True, color=INK)
        bullets(s, x + 0.2, 3.6, 2.55, 1.8, items, size=10.5, spacing=6, marker_color=c)
    rrect(s, 0.6, 5.7, 12.13, 1.1, fill=ACCENT_SOFT, line=None)
    text(s, 0.85, 5.8, 3, 0.3, "阶段 1 通过标准（Go / No-Go）", size=11.5, bold=True, color=ACCENT)
    text(s, 0.85, 6.12, 11.7, 0.6,
         "① 小艺 Tab 内连续预览 ≥ 30 min 无黑屏 / 崩溃；② 进出 Tab 主预览可恢复；③ meta 与帧 ts 对齐率接近 100%。"
         "若 ① 失败则切换到「宿主显示 + 小艺分析」架构，其余设计不变。", size=10.5, color=INK)

    # ------------------------------------------------------------ 11 结论与请求
    s = blank(prs)
    header(s, "结论与请求", "方案可行；请评审拍板会话归属、显示位置与拍摄路径三项", 11)
    concl = [
        ("架构", "相机 App 握相机 → 框架 → HDI → HAL，配「显示 + 分析」两路 Surface；小艺消费分析路做感知与推荐，不直连 HAL / HDI。", CAM),
        ("门控", "稳定帧门控独立成模块，挂 Tab 业务层；感知 / 模型 / AR / 深度都订同一事件。", XY),
        ("拍摄", "点击拍摄 = 按 timestamp 触发一次带深度的 Capture；RGB / 深度 / meta 同 ts 对齐。", DEP),
    ]
    for i, (t, d, c) in enumerate(concl):
        y = 1.45 + i * 1.05
        rrect(s, 0.6, y, 12.13, 0.85, fill=WHITE, line=LINE, line_w=0.75)
        rect(s, 0.6, y + 0.12, 0.07, 0.61, fill=c)
        pill(s, 0.9, y + 0.26, 0.9, 0.34, t, c, size=10.5)
        text(s, 2.0, y, 10.5, 0.85, d, size=13, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    rrect(s, 0.6, 4.75, 12.13, 2.05, fill=ACCENT, line=None)
    text(s, 0.95, 4.9, 4, 0.35, "请评审决策", size=13, bold=True, color=RGBColor(0xC7, 0xD2, 0xFE))
    asks = ["同意「相机握机、小艺消费」为基线架构，并授权与相机 / HAL 团队对齐接口契约",
            "同意先做阶段 1 样机预研（UEC 显示 + 会话共存 + meta 对齐）作为 Go / No-Go 门",
            "明确拍摄路径偏好：允许短暂停流取深度；RGB 用预览帧还是 Still"]
    for i, a in enumerate(asks):
        badge(s, 1.15, 5.5 + i * 0.42, i + 1, color=WHITE, r=0.14, size=9)
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT
        text(s, 1.45, 5.32 + i * 0.42, 11, 0.36, a, size=12, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    out_paths = [
        "/workspace/artifacts/小艺CameraTab_UEC感知方案_评审.pptx",
        "/opt/cursor/artifacts/小艺CameraTab_UEC感知方案_评审.pptx",
    ]
    for path in out_paths:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        prs.save(path)
        print("saved", path)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    build()
